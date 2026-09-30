import tempfile
from pathlib import Path
from unittest.mock import MagicMock
from openpyxl import load_workbook

from dto import AvitoConfig
from models import Item
from parser.export.excel import ExcelStorage as Excel
from parser_cls import AvitoParse


def test_extract_params_from_beduin_scenario():
    """Тест извлечения параметров недвижимости/товара из сценария Beduin (mobile card API)."""
    payload = {
        "success": {
            "view": {
                "scenario": {
                    "beduin": {
                        "main": {
                            "params": {
                                "itemParams": {
                                    "items": [
                                        {"title": "Общая площадь", "description": "120 м²"},
                                        {"title": "Вход", "description": "отдельный с улицы"},
                                        {"title": "Этаж", "description": "1 из 5"},
                                    ]
                                },
                                "aboutPremises": {
                                    "items": [
                                        {"title": "Отделка", "description": "офисная"},
                                        {"title": "Отопление", "description": "центральное"},
                                    ]
                                }
                            }
                        }
                    }
                }
            }
        }
    }

    params = AvitoParse._extract_params(payload)
    assert params is not None
    assert params["Общая площадь"] == "120 м²"
    assert params["Вход"] == "отдельный с улицы"
    assert params["Этаж"] == "1 из 5"
    assert params["Отделка"] == "офисная"
    assert params["Отопление"] == "центральное"


def test_extract_params_from_mobile_api():
    """Тест извлечения параметров из mobile.params / parameters."""
    payload = {
        "success": {
            "mobile": {
                "params": [
                    {"title": "Площадь", "description": "55 м²"},
                    {"title": "Тип здания", "description": "жилой дом"},
                ]
            }
        }
    }

    params = AvitoParse._extract_params(payload)
    assert params is not None
    assert params["Площадь"] == "55 м²"
    assert params["Тип здания"] == "жилой дом"


def test_extract_params_returns_none_when_empty():
    """Тест безопасной обработки пустых или некорректных ответов."""
    assert AvitoParse._extract_params({}) is None
    assert AvitoParse._extract_params({"success": {}}) is None
    assert AvitoParse._extract_params(None) is None


def test_parse_params_method_enriches_ad():
    """Тест метода parse_params при включённой опции в конфиге."""
    config = AvitoConfig(urls=["https://www.avito.ru/test"], parse_params=True)
    parser = AvitoParse.__new__(AvitoParse)
    parser.config = config
    parser.stop_event = None
    parser.good_request_count = 0
    parser.bad_request_count = 0
    parser.http = MagicMock()
    parser.http.fetch_item_data.return_value = {
        "success": {
            "view": {
                "scenario": {
                    "beduin": {
                        "main": {
                            "params": {
                                "itemParams": {
                                    "items": [
                                        {"title": "Вход", "description": "с улицы"},
                                        {"title": "Общая площадь", "description": "85.5 м²"},
                                    ]
                                }
                            }
                        }
                    }
                }
            }
        }
    }

    ad = Item(id=999, urlPath="/realty_999", title="Помещение свободного назначения")
    result = parser.parse_params([ad])

    assert len(result) == 1
    assert result[0].params is not None
    assert result[0].params["Вход"] == "с улицы"
    assert result[0].params["Общая площадь"] == "85.5 м²"
    parser.http.fetch_item_data.assert_called_once_with(999)


def test_excel_export_saves_params_column():
    """Тест выгрузки параметров объекта («О помещении») в XLSX."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "realty_export.xlsx"
        excel = Excel(file_path=file_path)

        ad = Item(
            id=777,
            title="Офис 100м²",
            params={
                "Вход": "отдельный",
                "Общая площадь": "100 м²",
                "Этаж": "2",
            }
        )

        excel.save([ad])

        wb = load_workbook(file_path)
        sheet = wb.active
        # Header check: last column should be "Параметры"
        headers = [cell.value for cell in sheet[1]]
        assert "Параметры" in headers
        params_col_idx = headers.index("Параметры") + 1

        # Row check
        cell_val = sheet.cell(row=2, column=params_col_idx).value
        assert "Вход: отдельный" in cell_val
        assert "Общая площадь: 100 м²" in cell_val
        assert "Этаж: 2" in cell_val


def test_excel_export_handles_none_params():
    """Тест выгрузки в XLSX при отсутствии параметров."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "empty_params_export.xlsx"
        excel = Excel(file_path=file_path)

        ad = Item(id=778, title="Товар без параметров", params=None)
        excel.save([ad])

        wb = load_workbook(file_path)
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]
        params_col_idx = headers.index("Параметры") + 1

        cell_val = sheet.cell(row=2, column=params_col_idx).value
        assert cell_val in ("", None)
